from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
p='Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx';d=Document(p)
for para in d.sections[0].first_page_footer.paragraphs:para.clear()
for para in d.sections[0].first_page_header.paragraphs:para.clear()
for para in d.paragraphs:
 if para.text=='MỤC LỤC':
  for x in para._p.xpath('.//w:br[@w:type="page"]'):x.getparent().remove(x)
 if para.text.startswith(('Miền:','Router ')):para.paragraph_format.keep_with_next=True
 if para.text.startswith('Quy tắc mức model:'):para.alignment=WD_ALIGN_PARAGRAPH.LEFT
 if re_match:=para.text.startswith('Em xin cam đoan đồ án'):
  for r in para.runs:
   if 'Phát triển hệ thống quản trị nhà hàng thông minh Crispy Bite' in r.text:r.text=r.text.replace('Phát triển hệ thống quản trị nhà hàng thông minh Crispy Bite','Phát triển hệ thống quản lý nhà hàng Crispy Bite')
 if para.text.startswith(tuple('['+str(i)+']' for i in range(1,10))):para.alignment=WD_ALIGN_PARAGRAPH.LEFT;para.paragraph_format.first_line_indent=Cm(0)
 if para.style.name=='Heading 41':
  pp=para._p.get_or_add_pPr();o=pp.find(qn('w:outlineLvl'))
  if o is None:o=OxmlElement('w:outlineLvl');pp.append(o)
  o.set(qn('w:val'),'3')
 if para.text.startswith('Các vị trí đã được ghi trực tiếp'):
  para.text='Bổ sung ảnh tại 11 vị trí ở Chương 4 theo bảng dưới đây. Dùng bản chạy thực tế và dữ liệu thử nghiệm; ghi nhận cả trạng thái thành công và lỗi có ý nghĩa.'
 if para.text.startswith('Chèn ảnh ngay phía trên Caption'):
  para.text='Chèn ảnh phía trên Caption tương ứng ở Chương 4; thay dòng ghi chú bằng ảnh thật và xóa cụm [chưa bổ sung ảnh]. Giữ trường SEQ trong Caption. Với hình mới, dùng References → Insert Caption và chọn nhãn Hình hoặc Bảng.'
  para.paragraph_format.line_spacing=1.15
 if para.text.startswith('Sau khi bổ sung, nhấn Ctrl+A'):
  para.text='Nhấn Ctrl+A → F9, chọn cập nhật toàn bộ mục lục rồi lưu. Các bảng đã có đầy đủ; cần bổ sung ảnh giao diện và ghi kết quả thực tế, log hoặc ảnh minh chứng cho các phiếu kiểm thử ở Phụ lục C.'
  para.paragraph_format.line_spacing=1.15
for ins in d.element.xpath('.//w:instrText'):
 if 'TOC ' in (ins.text or '') and '\\o "1-3"' in ins.text:ins.text=ins.text.replace('1-3','1-4')
if 'toc 4' not in d.styles:d.styles.add_style('toc 4',WD_STYLE_TYPE.PARAGRAPH)
s=d.styles['toc 4'];s.font.name='Times New Roman';s.font.size=Pt(12);s.paragraph_format.line_spacing=1.15;s.paragraph_format.left_indent=Cm(1.5);s.paragraph_format.space_after=Pt(3)
# Reduce an almost empty final list-of-tables page while keeping 12 pt type.
s=d.styles['table of figures'];s.paragraph_format.space_after=Pt(1);s.paragraph_format.line_spacing=1.12
for para in d.paragraphs:
 if para.style.name=='table of figures':para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.12
# Consistent explicit borders in every cell, including conditional table styles.
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   tcp=c._tc.get_or_add_tcPr();b=tcp.find(qn('w:tcBorders'))
   if b is not None:tcp.remove(b)
   b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side)
    for k,v in [('val','single'),('sz','4'),('color','999999')]:x.set(qn('w:'+k),v)
    b.append(x)
   tcp.append(b)
 h=[c.text for c in t.rows[0].cells]
 if h==['STT','Method','Đường dẫn API đầy đủ']:
  ws=[1.2,2.4,12.9]
  for col,w in zip(t._tbl.tblGrid.gridCol_lst,ws):col.set(qn('w:w'),str(round(w*567)))
  for row in t.rows:
   for c,w in zip(row.cells,ws):c.width=Cm(w)
 if h==['Module','Số route']:
  for row in t.rows[1:]:
   if row.cells[0].text=='employee-attendance' and row.cells[1].text=='1':row.cells[0].text='attendance-kiosk'
 if h==['Vị trí hình','Màn hình','Nội dung cần bổ sung','Trạng thái']:
  ws=[1.8,4.1,8.1,2.5]
  for col,w in zip(t._tbl.tblGrid.gridCol_lst,ws):col.set(qn('w:w'),str(round(w*567)))
  for ri,row in enumerate(t.rows):
   if ri:row.cells[3].text='Cần bổ sung'
   for c,w in zip(row.cells,ws):
    c.width=Cm(w)
    for para in c.paragraphs:
     para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
     for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5)
d.save(p);print('layout patched')
