from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from collections import defaultdict
from pathlib import Path
p=Path('Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx');d=Document(p)
def f(para,code,num):
 for typ in ['begin','code','separate','text','end']:
  r=OxmlElement('w:r')
  if typ in ['begin','separate','end']:x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),typ)
  elif typ=='code':x=OxmlElement('w:instrText');x.set('{http://www.w3.org/XML/1998/namespace}space','preserve');x.text=' '+code+' '
  else:x=OxmlElement('w:t');x.text=str(num)
  r.append(x);para._p.append(r)
import re
counts=defaultdict(int)
for para in d.paragraphs:
 if para.style.name!='Caption':continue
 m=re.match(r'^(Bảng|Hình) ([0-9A-D]+)\.(\d+)\s*:\s*(.*)$',para.text)
 if not m:continue
 kind,ch,old,title=m.groups();counts[(kind,ch)]+=1;idx=counts[(kind,ch)]
 para.clear();para.add_run(kind+' '+ch+'.');f(para,'SEQ '+kind+(' \\r 1' if idx==1 else ''),idx);para.add_run(': '+title)
 for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
d.save(p)
print('captions',sum(counts.values()))
