from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
import hashlib,json,re
root=Path.cwd();p=root/'Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx';d=Document(p);errors=[]
base=json.loads((root/'report_work_20261007/source_audit.json').read_text(encoding='utf-8'));ref=json.loads((root/'report_work_20261007/reference_audit.json').read_text(encoding='utf-8'))
hashes={}
for pat,audit in [('Tuan3*CanChinhSua.docx',base),('12523037*.docx',ref)]:
 f=next(root.glob(pat));h=hashlib.sha256(f.read_bytes()).hexdigest();hashes[f.name]=h==audit['sha256']
 if h!=audit['sha256']:errors.append('input changed '+f.name)
with ZipFile(p) as z:
 x=etree.fromstring(z.read('word/document.xml'));ns={'w':qn('w:t').split('}')[0][1:],'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 bms=set(x.xpath('//w:bookmarkStart/@w:name',namespaces=ns));anchors=x.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
 dangling=[a for a in anchors if a not in bms]
 if dangling:errors.append('dangling links '+str(len(dangling)))
 fldcount=len(x.xpath('//w:fldChar[@w:fldCharType="begin"]',namespaces=ns));endcount=len(x.xpath('//w:fldChar[@w:fldCharType="end"]',namespaces=ns))
 if fldcount!=endcount:errors.append('unbalanced fields')
 captions=[q for q in d.paragraphs if q.style.name=='Caption']
 missingseq=[q.text for q in captions if 'SEQ ' not in ''.join(q._p.xpath('.//w:instrText/text()'))]
 if missingseq:errors.append('missing SEQ '+str(len(missingseq)))
 for i,t in enumerate(d.tables,1):
  prev=t._tbl.getprevious();text=''.join(n.text or '' for n in prev.iter(qn('w:t')))
  if not text.startswith('Bảng '):errors.append('table lacks caption '+str(i))
 text='\n'.join(n.text or '' for n in x.iter(qn('w:t')))
 for forbidden in ['Error!','Lỗi!','Thiết kế website quản lý đào tạo','Đang cập nhật mục lục']:
  if forbidden in text:errors.append('unresolved '+forbidden)
 report={'file':str(p),'tables':len(d.tables),'pictures_excluding_logo':len(d.inline_shapes)-1,'captions':len(captions),'missing_image_notes':sum(q.text.startswith('[CẦN BỔ SUNG ẢNH CHỤP') for q in d.paragraphs),'fields_begin':fldcount,'fields_end':endcount,'hyperlink_anchors':len(anchors),'dangling_anchors':dangling,'inputs_unchanged':hashes,'errors':errors}
(root/'report_work_20261007/final_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
