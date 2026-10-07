from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from collections import Counter
import json
p=Path('Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx');d=Document(p)
issues=[]
for i,t in enumerate(d.tables,1):
 prev=t._tbl.getprevious();txt=''.join(x.text or '' for x in prev.iter(qn('w:t'))) if prev is not None else ''
 if not txt.startswith('Bảng '):issues.append(('table_missing_caption',i,txt[:60]))
for i,shape in enumerate(d.inline_shapes):
 par=shape._inline.getparent().getparent().getparent();following=par.getnext();txt=''.join(x.text or '' for x in following.iter(qn('w:t'))) if following is not None else ''
 if i>0 and not txt.startswith('Hình '):issues.append(('image_missing_caption',i,txt[:50]))
fields=[''.join(p._p.xpath('.//w:instrText/text()')) for p in d.paragraphs if p._p.xpath('.//w:instrText')]
print(json.dumps({'tables':len(d.tables),'images_with_cover':len(d.inline_shapes),'captions':sum(p.style.name=='Caption' for p in d.paragraphs),'TOCs':[x for x in fields if x.startswith(' TOC')],'SEQ':sum('SEQ' in x for x in fields),'issues':issues,'headings':Counter(p.style.name for p in d.paragraphs if p.style.name.startswith('Heading'))},ensure_ascii=False,indent=2))

