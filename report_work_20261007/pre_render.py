from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from pathlib import Path
p=Path('Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx');d=Document(p)
for t in d.tables:
 h=[c.text for c in t.rows[0].cells]
 if h==['STT','Method','Đường dẫn API đầy đủ']:
  ws=[1.2,2.0,13.3]
  for c,w in zip(t._tbl.tblGrid.gridCol_lst,ws):c.set(qn('w:w'),str(round(w*567)))
  for row in t.rows:
   for c,w in zip(row.cells,ws):c.width=Cm(w)
# Use exact explanatory term for diagrams with explicit FK targets.
for para in d.paragraphs:
 if para.style.name=='Caption' and 'ERD chi tiết' in para.text:
  for r in para.runs:r.text=r.text.replace('ERD chi tiết','Lược đồ quan hệ chi tiết')
 if para.text.startswith('Hướng dẫn bổ sung'):para.text='D.1. Hướng dẫn bổ sung và cập nhật tự động'
d.save(p)
print('ready')
