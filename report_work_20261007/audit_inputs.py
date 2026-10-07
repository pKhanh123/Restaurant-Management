from pathlib import Path
from docx import Document
from zipfile import ZipFile
from collections import Counter
import json, hashlib
root=Path.cwd(); out=root/'report_work_20261007'
for name,pattern in [('reference','12523037*.docx'),('source','Tuan3*CanChinhSua.docx')]:
 p=next(root.glob(pattern)); d=Document(p)
 lines=[]
 for i,e in enumerate(d.element.body):
  if e.tag.endswith('sectPr'): continue
  txt=''.join(e.itertext()) if False else ''.join([x.text or '' for x in e.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')])
  st=[x.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') for x in e.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')]
  if txt or list(e.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')): lines.append(f'{i:04d} [{st}] {txt}')
 (out/f'{name}_text.txt').write_text('\n'.join(lines),encoding='utf-8')
 info={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'pictures':len(d.inline_shapes),'sections':[{'size':[s.page_width.cm,s.page_height.cm],'margins':[s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm]} for s in d.sections], 'styles':[]}
 for sn in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2']:
  try:
   s=d.styles[sn]; pf=s.paragraph_format
   info['styles'].append({'name':sn,'font':s.font.name,'size':s.font.size.pt if s.font.size else None,'bold':s.font.bold,'alignment':str(pf.alignment),'spacing':str(pf.line_spacing),'before':str(pf.space_before),'after':str(pf.space_after)})
  except KeyError: pass
 with ZipFile(p) as z:
  info['fields']=d.element.xpath('.//w:instrText/text()')
  info['parts']=[{'name':x.filename,'size':x.file_size,'sha256':hashlib.sha256(z.read(x)).hexdigest()} for x in z.infolist()]
 (out/f'{name}_audit.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
 print(name, {k:v for k,v in info.items() if k not in ['parts','fields']})
 print('fields',info['fields'][:12])
 print('headings',[(p.style.name,p.text) for p in d.paragraphs if p.style.name.startswith('Heading')][:100])


