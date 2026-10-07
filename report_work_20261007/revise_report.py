from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from copy import deepcopy
from collections import Counter,defaultdict
from io import BytesIO
import re,json,hashlib
from PIL import Image,ImageDraw,ImageFont
ROOT=Path.cwd(); WORK=ROOT/'report_work_20261007'; OUT=ROOT/'Tuan3_PhanVanKhanh_BaoCaoDaChinhSua.docx'
ref=Document(next(ROOT.glob('12523037*.docx')));doc=Document(ROOT/'Tuan3_PhanVanKhanh_BaoCaoCanChinhSua.docx')
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'; R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def el(tag,**attrs):
 x=OxmlElement('w:'+tag)
 for k,v in attrs.items():x.set(qn('w:'+k),str(v))
 return x
def field(p,code,cached='1'):
 for typ in ['begin','code','separate','text','end']:
  r=OxmlElement('w:r')
  if typ in ['begin','separate','end']:x=el('fldChar',fldCharType=typ)
  elif typ=='code':x=el('instrText');x.set('{http://www.w3.org/XML/1998/namespace}space','preserve');x.text=' '+code+' '
  else:x=el('t');x.text=cached
  r.append(x);p._p.append(r)
def font_run(r,size=13,bold=None):
 r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
 rf=r._r.get_or_add_rPr().get_or_add_rFonts()
 for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
 for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  rf.attrib.pop(qn('w:'+a),None)
 if bold is not None:r.bold=bold

def insert_after(p,text,style='Nội dung'):
 e=OxmlElement('w:p');p._p.addnext(e);o=Paragraph(e,doc._body);o.style=style;o.add_run(text);return o

def settext(p,text):
 p.clear();p.add_run(text)

# Adopt reference cover as a complete component, preserving its logo relationships.
body=doc.element.body
for e in list(body)[:27]:body.remove(e)
cover=[]
for e in list(ref.element.body)[:21]:
 e=deepcopy(e)
 for n in e.iter():
  for a in list(n.attrib):
   if a.startswith('{'+R+'}') and n.attrib[a] in ref.part.rels:
    rel=ref.part.rels[n.attrib[a]]
    if rel.reltype.endswith('/image'):
     rid,_=doc.part.get_or_add_image(BytesIO(rel.target_part.blob));n.set(a,rid)
 for t in e.iter(qn('w:t')):
  if t.text and 'THIẾT KẾ WEBSITE QUẢN LÝ ĐÀO TẠO' in t.text:t.text='PHÁT TRIỂN HỆ THỐNG QUẢN LÝ NHÀ HÀNG CRISPY BITE'
 cover.append(e)
for e in reversed(cover):body.insert(0,e)
# Section geometry from reference; keep source's cover/body separation.
rs=ref.sections[0]
for si,s in enumerate(doc.sections):
 for name in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']:
  setattr(s,name,getattr(rs,name))
 s.different_first_page_header_footer=(si==0)
 if si==0:
  borders=s._sectPr.find(qn('w:pgBorders'))
  rb=rs._sectPr.find(qn('w:pgBorders'))
  if rb is not None:
   if borders is not None:s._sectPr.remove(borders)
   s._sectPr.append(deepcopy(rb))
 else:
  for b in s._sectPr.findall(qn('w:pgBorders')):s._sectPr.remove(b)
 for pn in s._sectPr.findall(qn('w:pgNumType')):pn.attrib.pop(qn('w:start'),None)

for name in ['Normal','Nội dung','Title','Heading 1','Heading 2','Heading 3','Heading 4','Caption','Report Item']:
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(13)
 s.font.bold=False;s.font.italic=False
 pf=s.paragraph_format;pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(3);pf.widow_control=True
 pf.keep_with_next=False;pf.page_break_before=False
 if name in ['Normal','Nội dung']:pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Cm(.65)
 if name.startswith('Heading'):
  s.font.bold=True;pf.keep_with_next=True;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.space_before=Pt(8)
  pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
  if name=='Heading 1':s.font.size=Pt(14);pf.page_break_before=True;pf.space_after=Pt(8)
 if name=='Caption':pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Cm(0);pf.line_spacing=1.15;pf.space_before=Pt(6);pf.space_after=Pt(6);s.font.size=Pt(12)
 if name=='Report Item':s.font.bold=True;pf.keep_with_next=True;pf.first_line_indent=Cm(0)
 if name=='Title':s.font.bold=True;s.font.size=Pt(14);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Cm(0);pf.line_spacing=1.15
for name in ['toc 1','toc 2','toc 3','table of figures']:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.color.rgb=RGBColor(0,0,0)
 pf=s.paragraph_format;pf.line_spacing=1.15;pf.space_before=Pt(0);pf.space_after=Pt(3);pf.first_line_indent=Cm(0);pf.keep_with_next=False
 if name.startswith('toc'):pf.left_indent=Cm((int(name[-1])-1)*.5);s.font.bold=name=='toc 1'
 else:pf.left_indent=Cm(0)
 # one right-aligned tab stop at text margin
 tabs=pf.tab_stops;tabs.clear_all()
 from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
 tabs.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

# Main chapters get genuine outline numbering.
numbering=doc.part.numbering_part.element
existing=[int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))];aid=max(existing+[0])+1
abstract=el('abstractNum',abstractNumId=aid);abstract.append(el('multiLevelType',val='multilevel'))
for lev,fmt in enumerate(['CHƯƠNG %1:','%1.%2.','%1.%2.%3.','%1.%2.%3.%4.']):
 lvl=el('lvl',ilvl=lev);lvl.append(el('start',val=1));lvl.append(el('numFmt',val='decimal'));lvl.append(el('lvlText',val=fmt));lvl.append(el('suff',val='space'));lvl.append(el('lvlJc',val='left'));abstract.append(lvl)
numbering.append(abstract);nid=max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]+[0])+1
num=el('num',numId=nid);num.append(el('abstractNumId',val=aid));numbering.append(num)
def setnum(p,level):
 pp=p._p.get_or_add_pPr();n=pp.find(qn('w:numPr'))
 if n is not None:pp.remove(n)
 n=el('numPr');n.append(el('ilvl',val=level));n.append(el('numId',val=nid));pp.append(n)

# Correct and enrich source prose using local implementation evidence.
replacements={
'Sản phẩm hướng đến môi trường học tập':'Sản phẩm hướng đến môi trường học tập và vận hành thử nghiệm. Phạm vi báo cáo gồm phân tích yêu cầu, thiết kế dữ liệu, cài đặt các phân hệ và xây dựng tiêu chí kiểm thử. Việc đánh giá tập trung vào tính đúng đắn của luồng đặt món, chế biến, thanh toán và dữ liệu quản trị.',
'Đối tượng là quy trình số hóa':'Đối tượng nghiên cứu là quy trình số hóa hoạt động nhà hàng phục vụ nhanh: đặt món, phục vụ tại bàn hoặc mang đi, điều phối bếp, thanh toán, tồn kho và quản trị nhân sự. Nhóm người dùng gồm khách hàng, thu ngân, nhân viên bếp và quản trị viên. Yêu cầu nghiệp vụ được xác định trong phạm vi xây dựng và thử nghiệm hệ thống Crispy Bite.',
'Những biểu đồ này':None,
'Các hình dưới đây mở rộng':'Các hình dưới đây phân rã lược đồ dữ liệu thành các nhóm thực thể. Mỗi ô thể hiện tên model và một số trường chính; dòng FK ghi rõ bảng đích theo khai báo @relation trong Prisma. Quan hệ được đọc từ trường FK sang khóa chính của bảng đích; dấu ? chỉ liên kết có thể rỗng. Phụ lục A cung cấp đầy đủ thuộc tính và ràng buộc.',
'Ngày 03/10/2026, lệnh kiểm tra kiểu':'Bản báo cáo ngày 03/10/2026 ghi nhận kiểm tra kiểu có mã thoát 0 và bộ kiểm thử frontend có 287/289 ca đạt. Đây là số liệu lịch sử được giữ lại để theo dõi; chưa có log chạy lại kèm bản biên tập ngày 07/10/2026 để xác nhận tình trạng hiện tại. Các kịch bản tích hợp và E2E ở Phụ lục C là tiêu chí nghiệm thu, không phải kết quả đã chạy.',
'Hai ca thất bại là':'Theo ghi nhận của bản nguồn ngày 03/10/2026, hai ca thất bại thuộc EmployeePayrollScreen.test.tsx và employeeScheduleScreen.test.tsx. Trước khi nghiệm thu cần chạy lại trên phiên bản mã nguồn đã chốt, lưu log và ghi kết quả thực tế; không dùng số liệu lịch sử để khẳng định toàn bộ hệ thống hiện đã đạt.',
'API hiện hành GET /api/reports/daily':'API GET /api/reports/daily dành cho ADMIN. Dịch vụ lọc đơn theo createdAt trong ngày Việt Nam (UTC+7), cộng finalAmount của đơn COMPLETED và tính giá trị đơn trung bình. Dashboard còn hiển thị món bán chạy, phương thức thanh toán và thời gian chuẩn bị. Báo cáo hiện chưa phải phân hệ quyết toán tài chính hoặc bộ báo cáo cuối ngày mở rộng.',
'Phụ lục này sinh từ':'Phụ lục mô tả các model trong backend/prisma/schema.prisma. Kiểu có dấu ? cho phép rỗng; [] là tập đối tượng liên quan ở lớp ORM, không phải cột vật lý trong MySQL. Chỉ các trường được khai báo trong @relation(fields: [...]) mới được xác định là khóa ngoại. Các mã tham chiếu nghiệp vụ như clientRequestId và sourceTransactionId không mặc nhiên là khóa ngoại. Ràng buộc @@ thể hiện chỉ mục hoặc quy tắc mức bảng.',
'Backend khai báo route theo module':'Backend khai báo route theo module trong app.ts. Middleware thực hiện xác thực, phân quyền; controller tiếp nhận đầu vào và service xử lý nghiệp vụ. Zod kiểm tra dữ liệu ở các luồng có schema. Phụ lục B ghi đầy đủ tiền tố API ghép từ app.ts và đường dẫn trong từng router để có thể đối chiếu trực tiếp khi triển khai.',
}
for p in list(doc.paragraphs):
 for prefix,new in replacements.items():
  if new and p.text.startswith(prefix):settext(p,new);break
 if 'ngày 03 tháng 10 năm 2026' in p.text:settext(p,p.text.replace('ngày 03 tháng 10 năm 2026','ngày 07 tháng 10 năm 2026'))
 if p.text.startswith('[4] Prisma,'):settext(p,'[4] Prisma, “Transactions”, https://www.prisma.io/docs/orm/fundamentals/transactions (truy cập 07/10/2026). Tham khảo nguyên lý giao dịch; phiên bản dự án được xác định từ package.json và package-lock.json.')
 elif re.match(r'^\[[1-8]\]',p.text):settext(p,p.text.replace('03/10/2026','07/10/2026'))
 if p.text.startswith('[9]'):settext(p,p.text.replace('03/10/2026','07/10/2026'))
 if p.text.startswith('FinancialAccount, CashVoucher'):
  insert_after(p,'Khi sổ quỹ được kích hoạt, thanh toán đơn tạo OrderPaymentTransaction và phiếu thu liên kết nguồn trong cùng giao dịch. Khoản cọc đã áp dụng được trừ khỏi số tiền phải thu; không tạo phiếu thu mới cho phần cọc đã thu trước. Phiếu hủy được đối ứng bằng chứng từ đảo và hệ thống kiểm tra số dư theo thời điểm giao dịch.')
 if p.text.startswith('API GET /api/reports/daily'):
  insert_after(p,'Giới hạn diễn giải: doanh thu hiện gắn với ngày tạo đơn, có thể khác ngày thu tiền. Phép tính grossProfit trong API lấy doanh thu trừ AUTO_DEDUCT và KITCHEN_WASTE; chưa bao gồm đầy đủ lương, chi phí vận hành và hoàn trả. Vì vậy chỉ số này không được dùng để kết luận lợi nhuận ròng hay thay thế đối soát sổ quỹ.')
 if p.text.startswith('TableOrderScreen nhận'):
  insert_after(p,'Đối với đơn QR gắn đặt bàn, máy chủ còn kiểm soát khai báo thanh toán, xác nhận của nhân viên và quyền trả sau. Khách không có quyền tự xác nhận tiền đã nhận hoặc gọi API thanh toán của thu ngân. Quyền khách hàng và phiên JWT của nhân viên được kiểm tra theo các luồng khác nhau.')

# Add detailed headings to navigation without treating captions as headings.
active=False;appx=None
for p in list(doc.paragraphs):
 t=p.text.strip();sn=p.style.name
 if t.startswith('CHƯƠNG 1:') and sn=='Heading 1':active=True
 if active and sn=='Heading 1' and not t.startswith('CHƯƠNG '):active=False
 if t.startswith('PHỤ LỤC '):appx=t[8:9]
 if sn=='Report Item':
  p.paragraph_format.page_break_before=False
  if re.match(r'^UC\d+:',t):p.style='Heading 3';settext(p,re.sub(r'^UC(\d+):\s*',lambda m:f'3.5.{int(m[1])}. UC{m[1]} — ',t))
  elif t.startswith('Miền '):p.style='Heading 4'
  elif re.match(r'^[AB]\.\d+',t):p.style='Heading 3' if appx=='A' else 'Heading 2'
  elif re.match(r'^TC\d+',t):p.style='Heading 2';settext(p,'C.'+str(int(t[2:4])+1)+'. '+t)
  elif appx=='A':p.style='Heading 2'
  elif t.startswith('Quy trình '):p.style='Heading 4'
 if active and p.style.name in ['Heading 1','Heading 2','Heading 3']:
  lev=int(p.style.name[-1])-1
  text=re.sub(r'^CHƯƠNG \d+:\s*','',p.text) if lev==0 else re.sub(r'^\d+(?:\.\d+)+\.\s*','',p.text)
  settext(p,text);setnum(p,lev)
 if t in ['NHẬN XÉT','LỜI CAM ĐOAN','LỜI CẢM ƠN']:
  p.style='Heading 1';p.paragraph_format.page_break_before=False
 if 'PHÁT TRIỂN HỆ THỐNG QUẢN LÝ NHÀ HÀNG' in t:p.style='Title'

# Normalize all non-cover body paragraphs; table styles handled separately.
cover_end=next(i for i,p in enumerate(doc.paragraphs) if p.text.strip()=='NHẬN XÉT')
for i,p in enumerate(doc.paragraphs):
 if i<cover_end:continue
 sn=p.style.name;pf=p.paragraph_format
 if sn.startswith('toc') or sn=='table of figures':continue
 if sn=='Caption':continue
 if p._p.xpath('.//w:drawing') or p._p.xpath('.//w:pict'):
  pf.keep_with_next=True;pf.first_line_indent=Cm(0);pf.line_spacing=1;pf.space_before=Pt(5);pf.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
 if sn.startswith('Heading'):
  pf.keep_with_next=True;pf.keep_together=True;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0)
  pf.line_spacing=1.3;pf.space_before=Pt(8);pf.space_after=Pt(5)
  if sn!='Heading 1':pf.page_break_before=False;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  for r in p.runs:font_run(r,14 if sn=='Heading 1' else 13,True)
 else:
  for r in p.runs:font_run(r,11 if p.text.startswith(('Miền:','Quy tắc mức model:')) else 13)
  if p.text.startswith(('Miền:','Quy tắc mức model:')):pf.line_spacing=1.15;pf.first_line_indent=Cm(0);pf.space_after=Pt(4)
  elif p.text.startswith('[CẦN BỔ SUNG'):
   pf.first_line_indent=Cm(0);pf.line_spacing=1.15;pf.space_before=Pt(8);pf.space_after=Pt(8);pf.keep_with_next=True
   for r in p.runs:r.italic=True;r.font.size=Pt(12)
  elif p.text.startswith('•'):pf.line_spacing=1.5
  elif sn=='Nội dung':pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(3)

# Remove cached TOC results as a complete field region, replace with genuine Word fields.
for title,code in [('MỤC LỤC','TOC \\o "1-3" \\h \\z \\u'),('DANH MỤC CÁC BẢNG','TOC \\h \\z \\c "Bảng"'),('DANH MỤC CÁC HÌNH VẼ, ĐỒ THỊ','TOC \\h \\z \\c "Hình"')]:
 p=next(p for p in doc.paragraphs if p.text.strip()==title)
 e=p._p.getnext()
 while e is not None:
  st=e.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
  if st is not None and st.get(qn('w:val'),'').startswith('Heading'):break
  nxt=e.getnext();body.remove(e);e=nxt
 q=insert_after(p,'','Normal');q.paragraph_format.first_line_indent=Cm(0);field(q,code,'Đang cập nhật mục lục')
 p.paragraph_format.page_break_before=True

# Rebuild captions with correct single SEQ instructions and visible missing-image flags.
counts=defaultdict(int)
for p in doc.paragraphs:
 if p.style.name!='Caption':continue
 m=re.match(r'^(Bảng|Hình) ([0-9A-D]+)\.(\d+)\s*:\s*(.*)$',p.text)
 if not m:continue
 kind,ch,old,title=m.groups();counts[(kind,ch)]+=1;idx=counts[(kind,ch)]
 p.clear();p.add_run(kind+' '+ch+'.');field(p,'SEQ '+kind+(' \\r 1' if idx==1 else ''),str(idx))
 if kind=='Hình' and ch=='4' and int(old)<=11:title+=' [chưa bổ sung ảnh]'
 p.add_run(': '+title);p.paragraph_format.keep_with_next=kind=='Bảng';p.paragraph_format.keep_together=True
 p.paragraph_format.page_break_before=False;p.paragraph_format.line_spacing=1.15;p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in p.runs:font_run(r,12,False)

# Parse current schema and repair dictionary meaning and false FK classification.
schema=(ROOT/'backend/prisma/schema.prisma').read_text(encoding='utf-8');models={}
for name,block in re.findall(r'(?ms)^model\s+(\w+)\s*\{(.*?)^\}',schema):
 fields=[]
 for line in block.splitlines():
  line=line.strip()
  if not line or line.startswith(('//','@@')):continue
  m=re.match(r'(\w+)\s+(\S+)(?:\s+(.*))?$',line)
  if m:fields.append((m[1],m[2],m[3] or ''))
 models[name]=fields
relations={}
for name,fs in models.items():
 relations[name]={}
 for f,typ,attrs in fs:
  m=re.search(r'@relation\(.*?fields:\s*\[([^]]+)\].*?references:\s*\[([^]]+)\]',attrs)
  if m:
   for key in m[1].split(','):relations[name][key.strip()]=(typ.rstrip('?[]'),m[2])
model_tables={}
for t in doc.tables:
 prev=t._tbl.getprevious();txt=''.join(prev.iter(qn('w:t'))).__str__() if False else ''.join(x.text or '' for x in prev.iter(qn('w:t'))) if prev is not None else ''
 if 'Cấu trúc model ' in txt:
  model=txt.split('Cấu trúc model ')[1];model_tables[model]=t
  for row in t.rows[1:]:
   if len(row.cells)<4:continue
   name=row.cells[0].text;typ=row.cells[1].text
   if name in relations.get(model,{}):
    target,key=relations[model][name];row.cells[2].text='Khóa ngoại đến '+target+'.'+key
   elif typ.rstrip('?[]') in models:row.cells[2].text=('Tập đối tượng liên quan ' if typ.endswith('[]') else 'Đối tượng liên quan ')+typ.rstrip('?[]')+'; trường ORM'
   elif 'Khóa ngoại' in row.cells[2].text:row.cells[2].text='Mã tham chiếu nghiệp vụ '+name+'; không khai báo FK'
   elif name=='requestDigest':row.cells[2].text='Mã băm nội dung yêu cầu'
   elif name=='idempotencyKey':row.cells[2].text='Khóa chống xử lý yêu cầu lặp'
   elif name=='sourceKey':row.cells[2].text='Khóa duy nhất của nghiệp vụ nguồn'
   elif name=='response':row.cells[2].text='Phản hồi lưu để trả lại khi gửi lặp'
# Honest historical results; accurate concrete endpoint examples.
for t in doc.tables:
 h=[c.text for c in t.rows[0].cells]
 if h==['Hạng mục','Kết quả','Nguồn bằng chứng']:
  vals=[['Kiểm tra kiểu TypeScript','Bản nguồn ghi đạt; chưa chạy lại','Ghi nhận 03/10/2026; cần đính kèm log'],['Frontend unit tests','Bản nguồn ghi 287/289 đạt','Ghi nhận 03/10/2026; cần chạy lại'],['Backend integration','Chưa có kết quả mới trong bản biên tập','Thực thi trên DB test cô lập'],['E2E và thiết bị','Chưa nghiệm thu','Bổ sung log và ảnh thực tế']]
  for row,vals2 in zip(t.rows[1:],vals):
   for c,v in zip(row.cells,vals2):c.text=v
 if h==['Chức năng','Method và path','Quyền']:
  vals=[['Báo cáo ngày','GET /api/reports/daily','ADMIN'],['Đăng nhập','POST /api/auth/login','Công khai'],['Danh sách món','GET /api/menu','Công khai'],['Tạo đơn','POST /api/orders','Theo phiên nhân viên hoặc token bàn/đặt bàn'],['Danh sách sổ quỹ','GET /api/cashbook','ADMIN, CASHIER']]
  for row,vs in zip(t.rows[1:],vals):
   for c,v in zip(row.cells,vs):c.text=v
# Full mount path for Appendix B.
app=(ROOT/'backend/src/app.ts').read_text(encoding='utf-8');mounts=dict((r,p) for p,r in re.findall(r"app\.use\('([^']+)',\s*(\w+)\)",app))
route_files={}
for f in (ROOT/'backend/src/modules').rglob('*.routes.ts'):
 txt=f.read_text(encoding='utf-8');m=re.search(r'export const (\w+)\s*=\s*Router',txt)
 if m:route_files[f.name]=(mounts.get(m[1],''),txt)
for t in doc.tables:
 prev=t._tbl.getprevious();txt=''.join(x.text or '' for x in prev.iter(qn('w:t'))) if prev is not None else ''
 if 'Các API của module ' not in txt:continue
 # Source note before caption contains filename.
 before=prev.getprevious();note=''.join(x.text or '' for x in before.iter(qn('w:t'))) if before is not None else ''
 match=next(((fn,pre) for fn,(pre,_) in route_files.items() if fn in note),None)
 if not match:continue
 fn,pre=match
 for row in t.rows[1:]:
  if len(row.cells)>1 and row.cells[1].text.startswith('/') and not row.cells[1].text.startswith('/api/'):
   row.cells[1].text=pre+row.cells[1].text.rstrip('/')
# Better missing-item index matching the 11 reserved figure captions.
last=doc.tables[-1]
rows=[('4.1','Đăng nhập và điều hướng ba vai trò','Đăng nhập; tab CASHIER, KITCHEN, ADMIN'),('4.2','POS','Giỏ hàng, tùy chọn món và xác nhận đơn'),('4.3','KDS','Vé chờ, đang chế biến, sẵn sàng'),('4.4','Gọi món QR','Thực đơn khách, giỏ hàng, trạng thái đơn'),('4.5','Bàn và đặt bàn','Sơ đồ bàn; danh sách, chi tiết đặt bàn'),('4.6','Thực đơn và giá','Danh mục, món, bảng giá, ưu đãi'),('4.7','Kho và định lượng','Nguyên liệu, BOM, phiếu nhập, kiểm kê'),('4.8','Khách hàng và giao hàng','Hồ sơ khách, nhóm khách, đối tác giao hàng'),('4.9','Nhân sự','Hồ sơ, lịch ca, công, lương, hoa hồng'),('4.10','Sổ quỹ','Tài khoản, phiếu thu, phiếu chi'),('4.11','Báo cáo và nhật ký','Dashboard ngày và thao tác kiểm toán')]
for row in list(last.rows)[1:]:last._tbl.remove(row._tr)
for c,v in zip(last.rows[0].cells,['Vị trí hình','Màn hình','Nội dung cần bổ sung','Trạng thái']):c.text=v
for vals in rows:
 for c,v in zip(last.add_row().cells,[*vals,'Chưa có ảnh chụp thực tế']):c.text=v
p=doc.add_paragraph('Hướng dẫn bổ sung và cập nhật tự động',style='Heading 2')
doc.add_paragraph('Chèn ảnh ngay phía trên Caption tương ứng ở Chương 4 và thay dòng ghi chú bằng ảnh thật. Giữ Caption có trường SEQ; xóa cụm [chưa bổ sung ảnh] sau khi đã chèn. Nếu cần thêm hình mới, dùng References → Insert Caption, chọn nhãn Hình hoặc Bảng để Word tiếp tục đánh số.',style='Nội dung')
doc.add_paragraph('Sau khi bổ sung, nhấn Ctrl+A rồi F9; với mục lục chọn cập nhật toàn bộ bảng. Cập nhật riêng danh mục hình và danh mục bảng nếu Word yêu cầu, sau đó lưu tài liệu. Các bảng đặc tả đã có đầy đủ; phần cần hoàn thiện là ảnh giao diện và kết quả thực tế/bằng chứng của các phiếu kiểm thử Phụ lục C.',style='Nội dung')

# Table layout preserves data but makes repeated headers and widths consistent.
for ti,t in enumerate(doc.tables):
 n=len(t.columns);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 h=[c.text for c in t.rows[0].cells]
 dictionary=h==['Trường','Kiểu','Ý nghĩa','Ràng buộc / mặc định']
 if dictionary:widths=[3.3,3.4,4.3,5.5]
 elif n==2:widths=[3.6,12.9]
 elif n==3:widths=[3.7,6.4,6.4]
 elif n==4:widths=[2.1,4.1,7,3.3]
 else:widths=[16.5/n]*n
 if h and h[0]=='Model nguồn':widths=[5.5,5.5,5.5]
 if h and h[0]=='Method':widths=[1.7,7.5,7.3] if n==3 else widths
 grid=t._tbl.tblGrid
 for col,w in zip(grid.gridCol_lst,widths):col.set(qn('w:w'),str(round(w*567)))
 pr=t._tbl.tblPr
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(round(16.5*567)));tw.set(qn('w:type'),'dxa')
 for old in pr.findall(qn('w:tblInd')):pr.remove(old)
 borders=pr.find(qn('w:tblBorders'))
 if borders is not None:pr.remove(borders)
 borders=el('tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:borders.append(el(side,val='single',sz=5,color='808080'))
 pr.append(borders)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  for x in trpr.findall(qn('w:trHeight')):trpr.remove(x)
  if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(el('tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcp=c._tc.get_or_add_tcPr()
   for tag in ['shd','tcBorders','tcMar','noWrap']:
    for x in tcp.findall(qn('w:'+tag)):tcp.remove(x)
   tcp.append(el('shd',fill='C5D9F1' if ri==0 else 'FFFFFF'))
   mar=el('tcMar')
   for side in ['top','bottom']:mar.append(el(side,w=65,type='dxa'))
   for side in ['left','right']:mar.append(el(side,w=75,type='dxa'))
   tcp.append(mar)
   for p in c.paragraphs:
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.line_spacing=1.12;pf.space_before=Pt(0);pf.space_after=Pt(2);pf.keep_with_next=ri==0;pf.keep_together=False;pf.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:font_run(r,10.5 if dictionary else 11,ri==0)

# All headers and footers use consistent page fields.
for si,s in enumerate(doc.sections):
 if si>0:s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
 for part in [s.header,s.footer]:
  for p in part.paragraphs:p.clear()
 hp=s.header.paragraphs[0];hp.text='Đồ án 3: Hệ thống quản lý nhà hàng Crispy Bite';hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.paragraph_format.first_line_indent=Cm(0);hp.paragraph_format.line_spacing=1
 for r in hp.runs:font_run(r,11)
 fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Cm(0);field(fp,'PAGE','1')
 for r in fp.runs:font_run(r,12)
# Set updating and custom caption labels.
settings=doc.settings.element
for u in settings.findall(qn('w:updateFields')):settings.remove(u)
settings.append(el('updateFields',val='true'))
caps=settings.find(qn('w:captions'))
if caps is None:caps=el('captions');settings.append(caps)
for name in ['Hình','Bảng']:
 if not any(x.get(qn('w:name'))==name for x in caps):caps.append(el('caption',name=name,pos='below' if name=='Hình' else 'above',numFmt='decimal'))
# Reassign drawing IDs after cover clone.
for i,x in enumerate(doc.element.xpath('.//wp:docPr'),1):x.set('id',str(i))
# Schema exact-coverage audit (not a claim of runtime behavior).
checks={'source_models':len(models),'document_models':len(model_tables),'missing_models':sorted(set(models)-set(model_tables)),'extra_models':sorted(set(model_tables)-set(models)),'mismatched_fields':{}}
for name,t in model_tables.items():
 actual=[(r.cells[0].text,r.cells[1].text) for r in t.rows[1:]];expected=[(a,b) for a,b,c in models.get(name,[])]
 if actual!=expected:checks['mismatched_fields'][name]={'actual':actual,'expected':expected}
(WORK/'schema_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
doc.core_properties.title='Phát triển hệ thống quản lý nhà hàng Crispy Bite';doc.core_properties.subject='Báo cáo Đồ án 3';doc.core_properties.author='Phan Văn Khánh';doc.core_properties.comments=''
doc.save(OUT)
print(OUT);print('tables',len(doc.tables),'images',len(doc.inline_shapes),'schema', {k:v for k,v in checks.items() if k!='mismatched_fields'},'mismatch',len(checks['mismatched_fields']))

