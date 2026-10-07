from pathlib import Path
from docx import Document
from collections import Counter
root=Path.cwd()
for pat in ['12523037*.docx','Tuan3*CanChinhSua.docx']:
 d=Document(next(root.glob(pat)))
 print('\nFILE',pat)
 print('styles',Counter(p.style.name for p in d.paragraphs))
 for name in ['Nội dung','Nidung','m1','m2','m3','CH1','ReportItem','Bìa']:
  for s in d.styles:
   if s.name==name or s.style_id==name:
    print(s.name,s.style_id,s.element.xml[:2800])
 print('captions',[(p.text, p._p.xpath('.//w:instrText/text()')) for p in d.paragraphs if p.style.name=='Caption'][:5])
 print('headers',[p.text for s in d.sections for p in s.header.paragraphs])
 print('footers',[p.text for s in d.sections for p in s.footer.paragraphs])
