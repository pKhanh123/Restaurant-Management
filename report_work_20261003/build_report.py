from __future__ import annotations

import os
import re
from pathlib import Path
from collections import defaultdict
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
WORK = Path(__file__).resolve().parent
SOURCE = Path(os.environ['TEMP']) / 'Tuan3_PhanVanKhanh_source.docx'
FINAL = ROOT / 'Tuan3_PhanVanKhanh_BaoCaoHoanChinh.docx'
DIAGRAMS = WORK / 'diagrams'
DIAGRAMS.mkdir(exist_ok=True)
SCHEMA = ROOT / 'backend' / 'prisma' / 'schema.prisma'

FONT_REG = r'C:\Windows\Fonts\times.ttf'
FONT_BOLD = r'C:\Windows\Fonts\timesbd.ttf'
FONT_ITALIC = r'C:\Windows\Fonts\timesi.ttf'

def font(size=30, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)

def wrap(draw, text, width, f):
    words = text.split()
    lines, cur = [], ''
    for word in words:
        test = (cur + ' ' + word).strip()
        if draw.textbbox((0, 0), test, font=f)[2] <= width:
            cur = test
        else:
            if cur: lines.append(cur)
            cur = word
    if cur: lines.append(cur)
    return lines

def center_text(draw, box, value, f, fill='#202020', line_gap=4):
    x0,y0,x1,y1 = box
    lines = wrap(draw, value, x1-x0-24, f)
    line_h = f.size + line_gap
    y = (y0+y1-len(lines)*line_h)/2
    for line in lines:
        w = draw.textbbox((0,0),line,font=f)[2]
        draw.text(((x0+x1-w)/2,y),line,font=f,fill=fill)
        y += line_h

def arrow(draw, a, b, color='#4E5B66', width=4):
    import math
    draw.line([a,b],fill=color,width=width)
    ang=math.atan2(b[1]-a[1],b[0]-a[0]); n=15
    left=(b[0]-n*math.cos(ang-0.55),b[1]-n*math.sin(ang-0.55))
    right=(b[0]-n*math.cos(ang+0.55),b[1]-n*math.sin(ang+0.55))
    draw.polygon([b,left,right],fill=color)

def base_canvas(title, subtitle='', w=1600, h=920):
    im=Image.new('RGB',(w,h),'white'); d=ImageDraw.Draw(im)
    d.text((55,38),title,font=font(39,True),fill='#1E2B36')
    if subtitle: d.text((56,90),subtitle,font=font(23),fill='#52616D')
    d.line((52,129,w-52,129),fill='#BAC6CE',width=3)
    return im,d

def save_diagram(im,name):
    path=DIAGRAMS/f'{name}.png'; im.save(path,optimize=True); return path

def draw_architecture():
    im,d=base_canvas('Kiến trúc tổng thể Crispy Bite','Các thành phần triển khai theo mã nguồn hiện tại')
    actors=[('Khách quét QR',65,205),('Thu ngân POS',65,375),('Bếp KDS',65,545),('Quản trị',65,715)]
    for label,x,y in actors:
        box=(x,y,x+275,y+95);d.rounded_rectangle(box,18,fill='#E8F0F4',outline='#536B7A',width=3);center_text(d,box,label,font(29,True))
    boxes=[('React Native / Expo Web',470,245,910,405),('Express REST API',1050,225,1515,385),('Socket.IO Gateway',1050,500,1515,660),('Prisma Client / MySQL',1050,710,1515,870)]
    for label,x0,y0,x1,y1 in boxes:
        box=(x0,y0,x1,y1);d.rounded_rectangle(box,22,fill='#F6F8FA',outline='#44515C',width=4);center_text(d,box,label,font(31,True))
    for _,x,y in actors: arrow(d,(x+275,y+47),(470,325))
    arrow(d,(910,300),(1050,305));arrow(d,(910,360),(1050,555));arrow(d,(1280,385),(1280,710));
    d.text((855,276),'HTTPS / JSON',font=font(22),fill='#4E5B66')
    d.text((895,465),'Sự kiện thời gian thực',font=font(21),fill='#4E5B66')
    return save_diagram(im,'architecture')

def draw_usecases(name,title,actor,items,others=None):
    im,d=base_canvas(title,'Sơ đồ ca sử dụng tổng hợp theo phân quyền và chức năng ứng dụng',h=1050)
    d.rounded_rectangle((355,170,1510,970),18,outline='#586B78',width=4)
    d.text((410,188),'HỆ THỐNG CRISPY BITE',font=font(29,True),fill='#394A55')
    d.ellipse((125,345,185,405),outline='#243746',width=5)
    d.line((155,405,155,530),fill='#243746',width=5)
    d.line((95,450,215,450),fill='#243746',width=5)
    d.line((155,530,100,610),fill='#243746',width=5)
    d.line((155,530,210,610),fill='#243746',width=5)
    center_text(d,(35,630,285,705),actor,font(31,True))
    n=len(items); cols=2 if n>6 else 1; rows=(n+cols-1)//cols
    positions=[]
    for j,item in enumerate(items):
        c=j%cols;r=j//cols
        x=430+c*535;y=250+r*(640/max(rows,1)); box=(x,y,x+435,y+86)
        d.ellipse(box,fill='#EFF4F7',outline='#4F6471',width=3);center_text(d,box,item,font(25))
        positions.append(((x+x+435)/2,(y+y+86)/2))
    for x,y in positions: d.line((230,465,x-220 if x>930 else x-220,y),fill='#9EABB4',width=2)
    if others:
        d.text((64,900),others,font=font(22),fill='#596875')
    return save_diagram(im,name)

def draw_erd(name,title,models,relations):
    im,d=base_canvas(title,'Ký hiệu PK là khóa chính; FK là khóa ngoại. Lược đồ rút gọn để dễ đọc.',h=1050)
    count=len(models); cols=3 if count>4 else 2; rows=(count+cols-1)//cols
    boxes={}
    for i,(model,fields) in enumerate(models):
        c=i%cols;r=i//cols; x=70+c*(1470/cols);y=180+r*(770/rows)
        box=(x,y,x+430,y+min(260,680/rows));boxes[model]=box
    for a,b in relations:
        if a in boxes and b in boxes:
            ax0,ay0,ax1,ay1=boxes[a];bx0,by0,bx1,by1=boxes[b]
            ac=((ax0+ax1)/2,(ay0+ay1)/2);bc=((bx0+bx1)/2,(by0+by1)/2)
            if abs(ac[1]-bc[1])<100:
                p=(ax1,ac[1]) if ac[0]<bc[0] else (ax0,ac[1])
                q=(bx0,bc[1]) if ac[0]<bc[0] else (bx1,bc[1])
            else:
                p=(ac[0],ay1) if ac[1]<bc[1] else (ac[0],ay0)
                q=(bc[0],by0) if ac[1]<bc[1] else (bc[0],by1)
            arrow(d,p,q,'#A0AFB9',3)
    for i,(model,fields) in enumerate(models):
        box=boxes[model];x,y,x1,y1=box
        d.rounded_rectangle(box,12,fill='#F7F9FA',outline='#617482',width=3)
        d.rounded_rectangle((x,y,x+430,y+52),12,fill='#DDE8EE',outline='#617482',width=3)
        heading_size=29 if len(model)<20 else 23
        d.text((x+14,y+10),model,font=font(heading_size,True),fill='#1A303E')
        yy=y+66
        for field in fields[:5]:
            for line in wrap(d,field,392,font(25))[:2]:
                d.text((x+18,yy),line,font=font(25),fill='#26343E');yy+=31
    return save_diagram(im,name)

def draw_sequence(name,title,lanes,steps):
    im,d=base_canvas(title,'Biểu đồ tuần tự rút gọn theo luồng nghiệp vụ',h=1150)
    xcoords=[130+i*(1350/(len(lanes)-1)) for i in range(len(lanes))]
    for x,label in zip(xcoords,lanes):
        box=(x-110,180,x+110,260);d.rounded_rectangle(box,12,fill='#E7EFF3',outline='#667986',width=3);center_text(d,box,label,font(24,True));d.line((x,260,x,1080),fill='#AFBAC1',width=2)
    for i,(a,b,label) in enumerate(steps):
        y=310+i*min(76,720/max(len(steps),1));x0=xcoords[a];x1=xcoords[b]
        arrow(d,(x0,y),(x1,y),'#526575',3)
        f=font(21);w=d.textbbox((0,0),label,font=f)[2]
        d.rectangle(((x0+x1-w)/2-8,y-31,(x0+x1+w)/2+8,y-3),fill='white')
        d.text(((x0+x1-w)/2,y-31),label,font=f,fill='#283944')
    return save_diagram(im,name)

def draw_flow(name,title,steps):
    im,d=base_canvas(title,'Các trạng thái và điểm kiểm soát chính',h=1120)
    n=len(steps);w=950;boxh=80;x=325
    for i,label in enumerate(steps):
        y=180+i*min(112,810/max(n-1,1));box=(x,y,x+w,y+boxh)
        fill='#E8F0F4' if i in (0,n-1) else '#F7F8F9'
        d.rounded_rectangle(box,18,fill=fill,outline='#5A6D7A',width=3)
        center_text(d,box,label,font(26,True if i in (0,n-1) else False))
        if i<n-1: arrow(d,(x+w/2,y+boxh),(x+w/2,y+boxh+30))
    return save_diagram(im,name)

def add_field(p, instruction, cached=''):
    r=p.add_run(); begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    r=p.add_run();instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+instruction+' ';r._r.append(instr)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    if cached:p.add_run(cached)
    r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)

def set_cell_border(cell,color='B8C1C8',size='5'):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders');tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        tag='w:'+edge;el=borders.find(qn(tag))
        if el is None:el=OxmlElement(tag);borders.append(el)
        el.set(qn('w:val'),'single');el.set(qn('w:sz'),size);el.set(qn('w:color'),color)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)

def set_cell_margin(cell,top=90,start=90,bottom=90,end=90):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();m=tcPr.first_child_found_in('w:tcMar')
    if m is None:m=OxmlElement('w:tcMar');tcPr.append(m)
    for key,val in (('top',top),('start',start),('bottom',bottom),('end',end)):
        el=m.find(qn('w:'+key))
        if el is None:el=OxmlElement('w:'+key);m.append(el)
        el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa')

class Report:
    def __init__(self):
        self.doc=Document(SOURCE)
        self._trim_source()
        self._styles()
        self.tables=defaultdict(int);self.figures=defaultdict(int)
        self.table_count=0;self.figure_count=0
        self._fix_front()

    def _trim_source(self):
        body=self.doc._element.body
        found=False
        for child in list(body):
            if child.tag==qn('w:p') and ''.join(child.itertext()).find('MỤC LỤC')>=0:
                found=True;continue
            if found and child.tag!=qn('w:sectPr'):body.remove(child)
        if not found:raise RuntimeError('Không tìm thấy mục lục của mẫu')

    def _styles(self):
        d=self.doc
        for name,size,bold,ital,align,before,after in [
            ('Heading 1',14,True,False,WD_ALIGN_PARAGRAPH.CENTER,18,12),
            ('Heading 2',13,True,False,WD_ALIGN_PARAGRAPH.JUSTIFY,13,5),
            ('Heading 3',13,True,True,WD_ALIGN_PARAGRAPH.JUSTIFY,9,3),
            ('Nội dung',13,False,False,WD_ALIGN_PARAGRAPH.JUSTIFY,4,4),
            ('Caption',11,False,True,WD_ALIGN_PARAGRAPH.CENTER,7,7),
        ]:
            s=d.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.italic=ital;s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.alignment=align;s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
            s.paragraph_format.line_spacing=1.5 if name=='Nội dung' else 1.15
            if name=='Nội dung':s.paragraph_format.first_line_indent=Cm(.65)
            if name.startswith('Heading'):s.paragraph_format.keep_with_next=True
        d.styles['Heading 1'].paragraph_format.page_break_before=True
        h1p=d.styles['Heading 1']._element.pPr
        if h1p is not None:
            num=h1p.find(qn('w:numPr'))
            if num is not None:h1p.remove(num)
        tof=d.styles['table of figures']
        tofp=tof._element.pPr
        if tofp is not None:
            num=tofp.find(qn('w:numPr'))
            if num is not None:tofp.remove(num)
        tof.font.name='Times New Roman';tof.font.size=Pt(11)
        tof.paragraph_format.left_indent=Cm(0);tof.paragraph_format.first_line_indent=Cm(0)
        tof.paragraph_format.space_before=Pt(1);tof.paragraph_format.space_after=Pt(0)
        tof.paragraph_format.line_spacing=1.05
        item=d.styles.add_style('Report Item',WD_STYLE_TYPE.PARAGRAPH)
        item.base_style=d.styles['Normal'];item.font.name='Times New Roman';item.font.size=Pt(13);item.font.bold=True
        item.font.color.rgb=RGBColor(0,0,0)
        item.paragraph_format.space_before=Pt(11);item.paragraph_format.space_after=Pt(5)
        item.paragraph_format.keep_with_next=True

    def _fix_front(self):
        ps=self.doc.paragraphs
        for p in ps:
            if 'HƯNG YÊN – 2025' in p.text:p.text=p.text.replace('2025','2026')
            if 'Crisby Bite' in p.text:p.text=p.text.replace('Crisby Bite','Crispy Bite')
            if 'ngày 13 tháng 09  năm 2026' in p.text:p.text=p.text.replace('ngày 13 tháng 09  năm 2026','ngày 03 tháng 10 năm 2026')
        for s in self.doc.sections[1:]:
            for p in s.header.paragraphs:
                if 'Đồ án 4:' in p.text:p.text=p.text.replace('Đồ án 4:','Đồ án 3:')

    def title(self,text):
        p=self.doc.add_paragraph(style='Heading 1');p.add_run(text);return p
    def h2(self,text):
        p=self.doc.add_paragraph(style='Heading 2');p.add_run(text);return p
    def h3(self,text):
        p=self.doc.add_paragraph(style='Heading 3');p.add_run(text);return p
    def item(self,text,break_before=False):
        p=self.doc.add_paragraph(style='Report Item');p.add_run(text)
        p.paragraph_format.page_break_before=break_before
        return p
    def body(self,text):
        p=self.doc.add_paragraph(style='Nội dung');p.add_run(text);return p
    def note(self,text):
        p=self.doc.add_paragraph(style='Normal');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1.1
        run=p.add_run(text);run.font.name='Times New Roman';run.font.size=Pt(10)
        return p
    def bullet(self,text):
        p=self.doc.add_paragraph(style='Nội dung');p.add_run('•  '+text)
        p.paragraph_format.left_indent=Cm(.6);p.paragraph_format.first_line_indent=Cm(-.35)
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
        return p
    def pagebreak(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def caption(self,kind,chapter,text):
        counts=self.tables if kind=='Bảng' else self.figures
        counts[chapter]+=1
        p=self.doc.add_paragraph(style='Caption')
        p.paragraph_format.keep_with_next = kind=='Bảng'
        p.add_run(f'{kind} {chapter}.')
        add_field(p,f'SEQ {kind} \\r 1' if counts[chapter]==1 else f'SEQ {kind}',str(counts[chapter]))
        p.add_run(f': {text}')
        if kind=='Bảng':self.table_count+=1
        else:self.figure_count+=1
        return p
    def table(self,chapter,title,headers,rows,widths=None,small=False):
        self.caption('Bảng',chapter,title)
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        if widths is None:widths=[16.1/len(headers)]*len(headers)
        for i,(cell,h) in enumerate(zip(t.rows[0].cells,headers)):
            cell.width=Cm(widths[i]);cell.text=str(h);shade(cell,'DDE8EE');set_cell_border(cell);set_cell_margin(cell)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9 if small else 10);r.font.bold=True
        trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trPr.append(repeat)
        for ri,row in enumerate(rows):
            cells=t.add_row().cells
            for i,(cell,value) in enumerate(zip(cells,row)):
                cell.width=Cm(widths[i]);cell.text=str(value);set_cell_border(cell);set_cell_margin(cell)
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri%2:shade(cell,'F8FAFB')
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8.5 if small else 9.5)
            cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
        return t
    def figure(self,chapter,title,path=None,note=None):
        if path:
            p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path),width=Cm(15.4))
        else:
            p=self.doc.add_paragraph(style='Nội dung')
            p.add_run('[CẦN BỔ SUNG ẢNH CHỤP: '+(note or title)+']').italic=True
        self.caption('Hình',chapter,title)
    def field_toc(self,kind):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
        instruction='TOC \\o "1-3" \\h \\z \\u' if kind=='toc' else f'TOC \\h \\z \\c "{kind}"'
        add_field(p,instruction,'(Cập nhật mục lục trong Word)')
    def save(self):
        settings=self.doc.settings.element
        upd=settings.find(qn('w:updateFields'))
        if upd is None:upd=OxmlElement('w:updateFields');settings.append(upd)
        upd.set(qn('w:val'),'true')
        self.doc.save(FINAL)
        return FINAL

def parse_schema():
    src=SCHEMA.read_text(encoding='utf-8')
    result=[]
    for name,body in re.findall(r'(?ms)^model\s+(\w+)\s*\{(.*?)^\}',src):
        fields=[];rules=[]
        for line in body.splitlines():
            s=line.strip()
            if not s or s.startswith('//'):continue
            if s.startswith('@@'):rules.append(s);continue
            parts=s.split(None,2)
            if len(parts)>=2:
                fields.append((parts[0],parts[1],parts[2] if len(parts)>2 else ''))
        result.append((name,fields,rules))
    return result

def parse_routes():
    results=[]
    for f in sorted((ROOT/'backend'/'src'/'modules').glob('**/*.routes.ts')):
        text=f.read_text(encoding='utf-8')
        found=re.findall(r'\.(get|post|put|patch|delete)\s*\(\s*[\'\"]([^\'\"]+)',text,re.I)
        results.append((f.parent.name,[(m.upper(),p) for m,p in found]))
    return results

if __name__=='__main__':
    from report_content import fill_report
    rep=Report();fill_report(rep)
    print(rep.save())
    print('tables',rep.table_count,'figures',rep.figure_count,'models',len(parse_schema()))
